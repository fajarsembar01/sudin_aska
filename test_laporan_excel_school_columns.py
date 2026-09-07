from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from dashboard.laporan import queries


def _worksheet_rows(xlsx_bytes: bytes, sheet_name: str) -> list[tuple]:
    workbook = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    return list(workbook[sheet_name].iter_rows(values_only=True))


def test_answer_export_includes_school_district_and_type(monkeypatch):
    monkeypatch.setattr(
        queries, "get_form", lambda _form_id: {"title": "Laporan Uji"}
    )
    monkeypatch.setattr(queries, "get_form_fields", lambda _form_id: [])
    monkeypatch.setattr(
        queries,
        "list_form_submissions",
        lambda _form_id: [
            {
                "id": 10,
                "school_id": 20,
                "status": "submitted",
                "school_name": "SD Contoh",
                "npsn": "12345678",
                "jenjang": "SD",
                "kecamatan_name": "CILINCING",
                "school_status": "SWASTA",
                "submitted_by_name": "Operator Sekolah",
            }
        ],
    )
    monkeypatch.setattr(
        queries, "get_submission_with_answers", lambda _submission_id: {"answers": []}
    )

    _filename, content = queries.export_form_xlsx(29)
    rows = _worksheet_rows(content, "Jawaban")

    assert rows[0][0:7] == (
        "No",
        "Sekolah",
        "NPSN",
        "Jenjang",
        "Kecamatan",
        "Jenis Sekolah",
        "Disubmit Oleh",
    )
    assert rows[1][0:7] == (
        1,
        "SD Contoh",
        "12345678",
        "SD",
        "CILINCING",
        "SWASTA",
        "Operator Sekolah",
    )


def test_no_submission_export_includes_school_district_and_type(monkeypatch):
    monkeypatch.setattr(
        queries, "get_form", lambda _form_id: {"title": "Laporan Uji"}
    )
    monkeypatch.setattr(
        queries,
        "get_form_target_schools",
        lambda _form: [
            {
                "id": 20,
                "name": "SD Contoh",
                "npsn": "12345678",
                "jenjang": "SD",
                "kecamatan_name": "KOJA",
                "status": "NEGERI",
            }
        ],
    )
    monkeypatch.setattr(queries, "list_form_submissions", lambda _form_id: [])

    _filename, content = queries.export_no_submissions_xlsx(29)
    rows = _worksheet_rows(content, "Tidak Mengumpulkan")

    assert rows[0] == (
        "No",
        "Sekolah",
        "NPSN",
        "Jenjang",
        "Kecamatan",
        "Jenis Sekolah",
        "Periode",
        "Status",
    )
    assert rows[1][0:6] == (
        1,
        "SD Contoh",
        "12345678",
        "SD",
        "KOJA",
        "NEGERI",
    )


def test_answer_export_filters_school_dimensions(monkeypatch):
    monkeypatch.setattr(
        queries, "get_form", lambda _form_id: {"title": "Laporan Uji"}
    )
    monkeypatch.setattr(queries, "get_form_fields", lambda _form_id: [])
    monkeypatch.setattr(
        queries,
        "list_form_submissions",
        lambda _form_id: [
            {
                "id": 10,
                "status": "submitted",
                "school_name": "SD Negeri",
                "jenjang": "SD",
                "kecamatan_name": "KOJA",
                "school_status": "NEGERI",
            },
            {
                "id": 11,
                "status": "submitted",
                "school_name": "SMP Swasta",
                "jenjang": "SMP",
                "kecamatan_name": "CILINCING",
                "school_status": "SWASTA",
            },
        ],
    )
    monkeypatch.setattr(
        queries, "get_submission_with_answers", lambda _submission_id: {"answers": []}
    )

    _filename, content = queries.export_form_xlsx(
        29, jenjang="smp", kecamatan="cilincing", school_status="swasta"
    )
    rows = _worksheet_rows(content, "Jawaban")

    assert len(rows) == 2
    assert rows[1][1] == "SMP Swasta"


def test_no_submission_export_filters_school_dimensions(monkeypatch):
    monkeypatch.setattr(
        queries, "get_form", lambda _form_id: {"title": "Laporan Uji"}
    )
    monkeypatch.setattr(
        queries,
        "get_form_target_schools",
        lambda _form: [
            {
                "id": 20,
                "name": "SD Negeri",
                "jenjang": "SD",
                "kecamatan_name": "KOJA",
                "status": "NEGERI",
            },
            {
                "id": 21,
                "name": "SMP Swasta",
                "jenjang": "SMP",
                "kecamatan_name": "CILINCING",
                "status": "SWASTA",
            },
        ],
    )
    monkeypatch.setattr(queries, "list_form_submissions", lambda _form_id: [])

    _filename, content = queries.export_no_submissions_xlsx(
        29, jenjang="SD", kecamatan="KOJA", school_status="NEGERI"
    )
    rows = _worksheet_rows(content, "Tidak Mengumpulkan")

    assert len(rows) == 2
    assert rows[1][1] == "SD Negeri"


def test_answer_export_accepts_multiple_jenjang_checkboxes(monkeypatch):
    monkeypatch.setattr(
        queries, "get_form", lambda _form_id: {"title": "Laporan Uji"}
    )
    monkeypatch.setattr(queries, "get_form_fields", lambda _form_id: [])
    monkeypatch.setattr(
        queries,
        "list_form_submissions",
        lambda _form_id: [
            {"id": 1, "status": "submitted", "school_name": "SD A", "jenjang": "SD"},
            {"id": 2, "status": "submitted", "school_name": "SMP B", "jenjang": "SMP"},
            {"id": 3, "status": "submitted", "school_name": "SMA C", "jenjang": "SMA"},
        ],
    )
    monkeypatch.setattr(
        queries, "get_submission_with_answers", lambda _submission_id: {"answers": []}
    )

    _filename, content = queries.export_form_xlsx(29, jenjang=["SD", "SMP"])
    rows = _worksheet_rows(content, "Jawaban")

    assert [row[1] for row in rows[1:]] == ["SD A", "SMP B"]
